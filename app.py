import json
import os
import sys
from datetime import date
from io import BytesIO
from getpass import getpass
from pathlib import Path
from urllib.parse import urlparse

from flask import (
    Flask,
    flash,
    jsonify,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from flask_login import (
    LoginManager,
    UserMixin,
    current_user,
    login_required,
    login_user,
    logout_user,
)
from flask_sqlalchemy import SQLAlchemy
from sqlalchemy import event, inspect, text
from werkzeug.security import check_password_hash, generate_password_hash


db = SQLAlchemy()
login_manager = LoginManager()

sys.modules.setdefault("app", sys.modules[__name__])


class Client(UserMixin, db.Model):
    __tablename__ = "clients"

    id = db.Column(db.Integer, primary_key=True)
    company_name = db.Column(db.String(150), nullable=False)
    email = db.Column(db.String(255), unique=True, nullable=False, index=True)
    password = db.Column(db.String(255), nullable=False)
    sharepoint_folder_path = db.Column(db.String(255), nullable=True)

    sales = db.relationship("Sales", back_populates="client", cascade="all, delete-orphan")
    purchases = db.relationship("Purchase", back_populates="client", cascade="all, delete-orphan")
    debtors = db.relationship("Debtor", back_populates="client", cascade="all, delete-orphan")
    profit_losses = db.relationship(
        "ProfitLoss", back_populates="client", cascade="all, delete-orphan"
    )

    def set_password(self, raw_password):
        self.password = generate_password_hash(raw_password)

    def check_password(self, raw_password):
        return check_password_hash(self.password, raw_password)


class Admin(db.Model):
    __tablename__ = "admins"

    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(150), nullable=False)
    email = db.Column(db.String(255), unique=True, nullable=False, index=True)
    password = db.Column(db.String(255), nullable=False)

    def set_password(self, raw_password):
        self.password = generate_password_hash(raw_password)

    def check_password(self, raw_password):
        return check_password_hash(self.password, raw_password)


class Sales(db.Model):
    __tablename__ = "sales"

    id = db.Column(db.Integer, primary_key=True)
    client_id = db.Column(db.Integer, db.ForeignKey("clients.id"), nullable=False, index=True)
    month = db.Column(db.String(20), nullable=False, index=True)
    amount = db.Column(db.Numeric(12, 2), nullable=False)

    client = db.relationship("Client", back_populates="sales")


class Purchase(db.Model):
    __tablename__ = "purchases"

    id = db.Column(db.Integer, primary_key=True)
    client_id = db.Column(db.Integer, db.ForeignKey("clients.id"), nullable=False, index=True)
    month = db.Column(db.String(20), nullable=False, index=True)
    amount = db.Column(db.Numeric(12, 2), nullable=False)

    client = db.relationship("Client", back_populates="purchases")


class Debtor(db.Model):
    __tablename__ = "debtors"

    id = db.Column(db.Integer, primary_key=True)
    client_id = db.Column(db.Integer, db.ForeignKey("clients.id"), nullable=False, index=True)
    party_name = db.Column(db.String(150), nullable=False)
    pending_amount = db.Column(db.Numeric(12, 2), nullable=False)
    month = db.Column(db.String(20), nullable=False, index=True)

    client = db.relationship("Client", back_populates="debtors")


class ProfitLoss(db.Model):
    __tablename__ = "profit_loss"

    id = db.Column(db.Integer, primary_key=True)
    client_id = db.Column(db.Integer, db.ForeignKey("clients.id"), nullable=False, index=True)
    month = db.Column(db.String(20), nullable=False, index=True)
    profit_loss_amount = db.Column(db.Numeric(12, 2), nullable=False)

    client = db.relationship("Client", back_populates="profit_losses")


def is_safe_next_url(target):
    if not target:
        return False

    parsed = urlparse(target)
    return parsed.scheme == "" and parsed.netloc == ""


MONTHS = [
    "April",
    "May",
    "June",
    "July",
    "August",
    "September",
    "October",
    "November",
    "December",
    "January",
    "February",
    "March",
]

SALES_INVOICE_COLUMNS = [
    ("date", "Date"),
    ("particulars", "Name"),
    ("gstin_uin", "GSTIN"),
    ("voucher_no.", "Invoice No."),
    ("sale", "Sale"),
    ("cgst", "CGST"),
    ("sgst", "SGST"),
    ("igst", "IGST"),
    ("round_off", "Round Off"),
    ("gross_total", "Invoice Value"),
]


def financial_year_for_date(value=None):
    value = value or date.today()
    start_year = value.year if value.month >= 4 else value.year - 1
    return f"{start_year}-{str(start_year + 1)[-2:]}"


def previous_financial_year(financial_year):
    start_year = int(str(financial_year).split("-", 1)[0])
    return f"{start_year - 1}-{str(start_year)[-2:]}"


def available_financial_years():
    current_financial_year = financial_year_for_date()
    return [previous_financial_year(current_financial_year), current_financial_year]


def selected_financial_year_from_request():
    financial_years = available_financial_years()
    selected_financial_year = request.args.get("financial_year", financial_years[-1])
    if selected_financial_year not in financial_years:
        selected_financial_year = financial_years[-1]
    return selected_financial_year

PURCHASE_INVOICE_COLUMNS = [
    ("date", "Date"),
    ("particulars", "Name"),
    ("gstin_uin", "GSTIN"),
    ("voucher_no.", "Invoice No."),
    ("purchase", "Purchase"),
    ("cgst", "CGST"),
    ("sgst", "SGST"),
    ("igst", "IGST"),
    ("round_off", "Round Off"),
    ("gross_total", "Invoice Value"),
]


def previous_month_name(month):
    month_index = MONTHS.index(month)
    return MONTHS[month_index - 1] if month_index > 0 else MONTHS[-1]


def parse_number(value):
    if value is None:
        return 0
    try:
        return float(str(value).replace(",", "") or 0)
    except ValueError:
        return 0


def money_total(connection, table_name, amount_column, client_id, month, financial_year):
    cursor = connection.execute(
        f"""
        SELECT COALESCE(SUM({amount_column}), 0)
        FROM {table_name}
        WHERE client_id = ? AND month = ? AND financial_year = ?
        """,
        (client_id, month, financial_year),
    )
    return float(cursor.fetchone()[0] or 0)


def row_count(connection, table_name, client_id, month, financial_year):
    cursor = connection.execute(
        f"""
        SELECT COUNT(*)
        FROM {table_name}
        WHERE client_id = ? AND month = ? AND financial_year = ?
        """,
        (client_id, month, financial_year),
    )
    return int(cursor.fetchone()[0] or 0)


def get_dashboard_data(client_id, month, financial_year=None):
    financial_year = financial_year or financial_year_for_date()
    from excel_upload import connect_upload_database, init_upload_database

    init_upload_database()
    with connect_upload_database() as connection:
        total_sales = money_total(connection, "sales", "amount", client_id, month, financial_year)
        total_purchase = money_total(connection, "purchase", "amount", client_id, month, financial_year)
        total_debtors = money_total(connection, "debtors", "pending_amount", client_id, month, financial_year)
        uploaded_profit_loss = money_total(
            connection, "profit_loss", "profit_loss_amount", client_id, month, financial_year
        )
        profit_loss_count = row_count(connection, "profit_loss", client_id, month, financial_year)
        profit_loss = uploaded_profit_loss if profit_loss_count else total_sales - total_purchase

        debtor_rows = connection.execute(
            """
            SELECT party_name, pending_amount
            FROM debtors
            WHERE client_id = ? AND month = ? AND financial_year = ?
            ORDER BY party_name
            """,
            (client_id, month, financial_year),
        ).fetchall()

    return {
        "month": month,
        "financial_year": financial_year,
        "total_sales": total_sales,
        "total_purchase": total_purchase,
        "profit_loss": profit_loss,
        "total_debtors": total_debtors,
        "debtors": [
            {"party_name": party_name, "pending_amount": float(pending_amount)}
            for party_name, pending_amount in debtor_rows
        ],
    }


def get_sales_invoice_details(client_id, month, financial_year=None):
    financial_year = financial_year or financial_year_for_date()
    from excel_upload import connect_upload_database, init_upload_database

    init_upload_database()
    with connect_upload_database() as connection:
        rows = connection.execute(
            """
            SELECT id, row_data
            FROM sales_invoices
            WHERE client_id = ? AND month = ? AND financial_year = ?
            ORDER BY id
            """,
            (client_id, month, financial_year),
        ).fetchall()

    records = []
    for invoice_id, row_data in rows:
        source_record = json.loads(row_data)
        record = {"_invoice_id": invoice_id}
        for key, label in SALES_INVOICE_COLUMNS:
            value = source_record.get(key, "")
            if key == "date" and isinstance(value, str) and "T" in value:
                value = value.split("T", 1)[0]
            record[label] = "" if value is None else value
        records.append(record)

    return records, [label for _, label in SALES_INVOICE_COLUMNS]


def get_purchase_invoice_details(client_id, month, financial_year=None):
    financial_year = financial_year or financial_year_for_date()
    from excel_upload import connect_upload_database, init_upload_database

    init_upload_database()
    with connect_upload_database() as connection:
        rows = connection.execute(
            """
            SELECT id, row_data
            FROM purchase_invoices
            WHERE client_id = ? AND month = ? AND financial_year = ?
            ORDER BY id
            """,
            (client_id, month, financial_year),
        ).fetchall()

    records = []
    for invoice_id, row_data in rows:
        source_record = json.loads(row_data)
        record = {"_invoice_id": invoice_id}
        for key, label in PURCHASE_INVOICE_COLUMNS:
            value = source_record.get(key, "")
            if key == "date" and isinstance(value, str) and "T" in value:
                value = value.split("T", 1)[0]
            record[label] = "" if value is None else value
        records.append(record)

    return records, [label for _, label in PURCHASE_INVOICE_COLUMNS]


def get_sales_invoice_rows(client_id, month, financial_year=None):
    financial_year = financial_year or financial_year_for_date()
    from excel_upload import connect_upload_database, init_upload_database

    init_upload_database()
    with connect_upload_database() as connection:
        rows = connection.execute(
            """
            SELECT row_data
            FROM sales_invoices
            WHERE client_id = ? AND month = ? AND financial_year = ?
            ORDER BY id
            """,
            (client_id, month, financial_year),
        ).fetchall()

    return [json.loads(row_data) for (row_data,) in rows]


def summarize_sales_rows(rows):
    customer_totals = {}
    total_sales = 0
    total_cgst = 0
    total_sgst = 0
    total_igst = 0
    total_round_off = 0
    total_invoice_value = 0

    for row in rows:
        sale = parse_number(row.get("sale"))
        cgst = parse_number(row.get("cgst"))
        sgst = parse_number(row.get("sgst"))
        igst = parse_number(row.get("igst"))
        round_off = parse_number(row.get("round_off"))
        invoice_value = parse_number(row.get("gross_total"))
        customer_name = str(row.get("particulars") or "Unknown party").strip() or "Unknown party"

        total_sales += sale
        total_cgst += cgst
        total_sgst += sgst
        total_igst += igst
        total_round_off += round_off
        total_invoice_value += invoice_value
        customer_totals[customer_name] = customer_totals.get(customer_name, 0) + sale

    invoice_count = len(rows)
    average_invoice_value = total_invoice_value / invoice_count if invoice_count else 0
    top_customers = sorted(
        [{"name": name, "amount": amount} for name, amount in customer_totals.items()],
        key=lambda item: item["amount"],
        reverse=True,
    )[:5]

    return {
        "total_sales": total_sales,
        "invoice_count": invoice_count,
        "average_invoice_value": average_invoice_value,
        "total_cgst": total_cgst,
        "total_sgst": total_sgst,
        "total_igst": total_igst,
        "total_round_off": total_round_off,
        "total_tax": total_cgst + total_sgst + total_igst,
        "top_customers": top_customers,
    }


def percent_change(current_value, previous_value):
    if previous_value == 0:
        return None if current_value else 0
    return ((current_value - previous_value) / previous_value) * 100


def format_percent(value):
    if value is None:
        return "N/A"
    return f"{value:.1f}%"


def build_sales_insights(current_summary, previous_summary, month, previous_month):
    sales_change = current_summary["total_sales"] - previous_summary["total_sales"]
    sales_change_percent = percent_change(
        current_summary["total_sales"],
        previous_summary["total_sales"],
    )
    invoice_change = current_summary["invoice_count"] - previous_summary["invoice_count"]
    average_change = (
        current_summary["average_invoice_value"] - previous_summary["average_invoice_value"]
    )

    if current_summary["invoice_count"] == 0:
        return f"No sales invoices were found for {month}. Import or sync invoices to generate analysis."

    if previous_summary["invoice_count"] == 0:
        return (
            f"{month} has {current_summary['invoice_count']} invoices and no invoice data was found "
            f"for {previous_month}, so this is the first comparable month in the portal."
        )

    direction = "increased" if sales_change >= 0 else "decreased"
    invoice_direction = "more" if invoice_change >= 0 else "fewer"
    average_direction = "higher" if average_change >= 0 else "lower"
    driver = (
        "higher invoice volume"
        if abs(invoice_change) > 0 and abs(average_change) <= abs(sales_change)
        else "a higher average invoice value"
        if average_change > 0
        else "a lower average invoice value"
    )

    return (
        f"Sales {direction} by {format_percent(abs(sales_change_percent))} compared with "
        f"{previous_month}. Invoice count is {abs(invoice_change)} {invoice_direction} than last month, "
        f"and average invoice value is {average_direction}. The main movement appears driven by {driver}."
    )


def get_sales_analysis_data(client_id, month, financial_year):
    previous_month = previous_month_name(month)
    previous_financial_year_value = (
        previous_financial_year(financial_year) if month == "April" else financial_year
    )
    current_summary = summarize_sales_rows(get_sales_invoice_rows(client_id, month, financial_year))
    previous_summary = summarize_sales_rows(
        get_sales_invoice_rows(client_id, previous_month, previous_financial_year_value)
    )

    sales_change = current_summary["total_sales"] - previous_summary["total_sales"]
    invoice_count_change = current_summary["invoice_count"] - previous_summary["invoice_count"]
    average_invoice_change = (
        current_summary["average_invoice_value"] - previous_summary["average_invoice_value"]
    )

    return {
        "month": month,
        "previous_month": previous_month,
        "current": current_summary,
        "previous": previous_summary,
        "sales_change": sales_change,
        "sales_change_percent": percent_change(
            current_summary["total_sales"],
            previous_summary["total_sales"],
        ),
        "invoice_count_change": invoice_count_change,
        "average_invoice_change": average_invoice_change,
        "average_invoice_change_percent": percent_change(
            current_summary["average_invoice_value"],
            previous_summary["average_invoice_value"],
        ),
        "insight_summary": build_sales_insights(
            current_summary,
            previous_summary,
            month,
            previous_month,
        ),
    }


def recalculate_sales_total(connection, client_id, month):
    rows = connection.execute(
        """
        SELECT row_data
        FROM sales_invoices
        WHERE client_id = ? AND month = ?
        """,
        (client_id, month),
    ).fetchall()

    total_sales = 0
    for (row_data,) in rows:
        value = json.loads(row_data).get("sale", 0)
        try:
            total_sales += float(value or 0)
        except (TypeError, ValueError):
            continue

    connection.execute(
        """
        DELETE FROM sales
        WHERE client_id = ? AND month = ?
        """,
        (client_id, month),
    )

    if total_sales:
        connection.execute(
            """
            INSERT INTO sales (client_id, month, amount)
            VALUES (?, ?, ?)
            """,
            (client_id, month, total_sales),
        )


def delete_sales_invoice_rows(invoice_ids, client_id, month):
    from excel_upload import connect_upload_database, init_upload_database

    invoice_ids = [int(invoice_id) for invoice_id in invoice_ids if str(invoice_id).isdigit()]
    if not invoice_ids:
        return 0

    init_upload_database()
    placeholders = ",".join("?" for _ in invoice_ids)
    with connect_upload_database() as connection:
        existing_rows = connection.execute(
            f"""
            SELECT id
            FROM sales_invoices
            WHERE client_id = ? AND month = ? AND id IN ({placeholders})
            """,
            [client_id, month, *invoice_ids],
        ).fetchall()

        existing_ids = [row[0] for row in existing_rows]
        if not existing_ids:
            return 0

        delete_placeholders = ",".join("?" for _ in existing_ids)
        connection.execute(
            f"""
            DELETE FROM sales_invoices
            WHERE client_id = ? AND month = ? AND id IN ({delete_placeholders})
            """,
            [client_id, month, *existing_ids],
        )
        recalculate_sales_total(connection, client_id, month)

    return len(existing_ids)


def make_invoice_excel(records, columns, sheet_title):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_title[:31]
    worksheet.append(columns)

    for cell in worksheet[1]:
        cell.font = Font(bold=True)

    for record in records:
        worksheet.append([record.get(column, "") for column in columns])

    for column_cells in worksheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(max_length + 2, 32)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def make_sales_excel(records, columns, month):
    return make_invoice_excel(records, columns, f"{month} Sales")


def pdf_escape(value):
    text = str(value or "").replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    return text.encode("latin-1", "replace").decode("latin-1")


def make_simple_pdf(title, records, columns):
    lines = [title, ""]
    lines.append(" | ".join(columns))
    lines.append("-" * 120)

    for record in records:
        row_values = [str(record.get(column, ""))[:24] for column in columns]
        lines.append(" | ".join(row_values))

    pages = [lines[index : index + 32] for index in range(0, len(lines), 32)] or [[]]
    objects = []
    page_ids = []
    font_id = 3

    def add_object(content):
        objects.append(content)
        return len(objects)

    add_object("<< /Type /Catalog /Pages 2 0 R >>")
    add_object("")
    add_object("<< /Type /Font /Subtype /Type1 /BaseFont /Courier >>")

    for page_lines in pages:
        stream_lines = ["BT", "/F1 8 Tf", "36 560 Td", "12 TL"]
        for line in page_lines:
            stream_lines.append(f"({pdf_escape(line)}) Tj")
            stream_lines.append("T*")
        stream_lines.append("ET")
        stream = "\n".join(stream_lines)
        content_id = add_object(f"<< /Length {len(stream.encode('latin-1'))} >>\nstream\n{stream}\nendstream")
        page_id = add_object(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 842 595] "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >>"
        )
        page_ids.append(page_id)

    objects[1] = f"<< /Type /Pages /Kids [{' '.join(f'{page_id} 0 R' for page_id in page_ids)}] /Count {len(page_ids)} >>"

    pdf = BytesIO()
    pdf.write(b"%PDF-1.4\n")
    offsets = [0]
    for object_id, content in enumerate(objects, start=1):
        offsets.append(pdf.tell())
        pdf.write(f"{object_id} 0 obj\n{content}\nendobj\n".encode("latin-1"))

    xref_position = pdf.tell()
    pdf.write(f"xref\n0 {len(objects) + 1}\n".encode("latin-1"))
    pdf.write(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.write(f"{offset:010d} 00000 n \n".encode("latin-1"))
    pdf.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_position}\n%%EOF".encode(
            "latin-1"
        )
    )
    pdf.seek(0)
    return pdf


def create_app():
    app = Flask(__name__)
    app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY", "change-this-secret-key")
    default_database_path = Path(app.root_path).joinpath("client_portal_main.sqlite3")
    app.config["SQLALCHEMY_DATABASE_URI"] = os.environ.get(
        "DATABASE_URL",
        f"sqlite:///{default_database_path.as_posix()}",
    )
    app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
    app.config["SQLALCHEMY_ENGINE_OPTIONS"] = {
        "connect_args": {"timeout": 30},
        "pool_pre_ping": True,
    }
    app.config["UPLOAD_SQLITE_DB"] = os.environ.get("UPLOAD_SQLITE_DB", "uploads.sqlite3")
    db.init_app(app)
    with app.app_context():
        engine = db.engine

    @event.listens_for(engine, "connect")
    def configure_sqlite_connection(dbapi_connection, connection_record):
        cursor = dbapi_connection.cursor()
        cursor.execute("PRAGMA journal_mode=MEMORY")
        cursor.execute("PRAGMA temp_store=MEMORY")
        cursor.execute("PRAGMA busy_timeout=30000")
        cursor.close()

    login_manager.init_app(app)
    login_manager.login_view = "login"
    login_manager.login_message_category = "warning"

    from admin_panel import admin_bp
    from excel_upload import excel_upload_bp, init_upload_database

    app.register_blueprint(admin_bp)
    app.register_blueprint(excel_upload_bp)

    with app.app_context():
        db.create_all()
        client_columns = {column["name"] for column in inspect(db.engine).get_columns("clients")}
        if "sharepoint_folder_path" not in client_columns:
            with db.engine.begin() as connection:
                connection.execute(
                    text("ALTER TABLE clients ADD COLUMN sharepoint_folder_path VARCHAR(255)")
                )

    @login_manager.user_loader
    def load_user(client_id):
        return db.session.get(Client, int(client_id))

    @app.route("/")
    def index():
        if current_user.is_authenticated:
            return redirect(url_for("dashboard"))
        return redirect(url_for("login"))

    @app.route("/login", methods=["GET", "POST"])
    def login():
        if current_user.is_authenticated:
            return redirect(url_for("dashboard"))

        if request.method == "POST":
            email = request.form.get("email", "").strip().lower()
            password = request.form.get("password", "")
            remember = request.form.get("remember") == "on"

            client = Client.query.filter_by(email=email).first()
            if client and client.check_password(password):
                login_user(client, remember=remember)
                next_page = request.args.get("next")
                if is_safe_next_url(next_page):
                    return redirect(next_page)
                return redirect(url_for("dashboard"))

            flash("Invalid email or password.", "danger")

        return render_template("login.html")

    @app.route("/logout")
    @login_required
    def logout():
        logout_user()
        flash("You have been logged out.", "info")
        return redirect(url_for("login"))

    @app.route("/admin/logout")
    def admin_logout():
        session.pop("admin_id", None)
        session.pop("admin_name", None)
        flash("Admin has been logged out.", "info")
        return redirect(url_for("admin.login"))

    @app.route("/dashboard")
    @login_required
    def dashboard():
        selected_month = request.args.get("month", "April")
        if selected_month not in MONTHS:
            selected_month = "April"
        financial_years = available_financial_years()
        selected_financial_year = selected_financial_year_from_request()

        return render_template(
            "dashboard.html",
            months=MONTHS,
            financial_years=financial_years,
            selected_financial_year=selected_financial_year,
            selected_month=selected_month,
            dashboard_data=get_dashboard_data(current_user.id, selected_month, selected_financial_year),
        )

    @app.route("/dashboard/data")
    @login_required
    def dashboard_data():
        selected_month = request.args.get("month", "April")
        if selected_month not in MONTHS:
            return jsonify({"error": "Invalid month selected."}), 400
        selected_financial_year = selected_financial_year_from_request()

        return jsonify(get_dashboard_data(current_user.id, selected_month, selected_financial_year))

    @app.route("/dashboard/sales")
    @login_required
    def sales_details():
        selected_month = request.args.get("month", "April")
        if selected_month not in MONTHS:
            selected_month = "April"
        financial_years = available_financial_years()
        selected_financial_year = selected_financial_year_from_request()

        records, columns = get_sales_invoice_details(
            current_user.id,
            selected_month,
            selected_financial_year,
        )
        return render_template(
            "sales_details.html",
            detail_type="sales",
            page_title="Sales Invoices",
            list_title="Sales invoice list",
            empty_message="No sales invoice details found",
            financial_years=financial_years,
            selected_financial_year=selected_financial_year,
            admin_view=False,
            months=MONTHS,
            selected_month=selected_month,
            records=records,
            columns=columns,
        )

    @app.route("/dashboard/purchase")
    @login_required
    def purchase_details():
        selected_month = request.args.get("month", "April")
        if selected_month not in MONTHS:
            selected_month = "April"
        financial_years = available_financial_years()
        selected_financial_year = selected_financial_year_from_request()

        records, columns = get_purchase_invoice_details(
            current_user.id,
            selected_month,
            selected_financial_year,
        )
        return render_template(
            "sales_details.html",
            detail_type="purchase",
            page_title="Purchase Invoices",
            list_title="Purchase invoice list",
            empty_message="No purchase invoice details found",
            financial_years=financial_years,
            selected_financial_year=selected_financial_year,
            admin_view=False,
            months=MONTHS,
            selected_month=selected_month,
            records=records,
            columns=columns,
        )

    @app.route("/dashboard/sales/analysis")
    @login_required
    def sales_analysis():
        selected_month = request.args.get("month", "April")
        if selected_month not in MONTHS:
            selected_month = "April"
        financial_years = available_financial_years()
        selected_financial_year = selected_financial_year_from_request()

        return render_template(
            "sales_analysis.html",
            months=MONTHS,
            financial_years=financial_years,
            selected_financial_year=selected_financial_year,
            selected_month=selected_month,
            analysis=get_sales_analysis_data(current_user.id, selected_month, selected_financial_year),
        )

    @app.route("/dashboard/sales/export/excel")
    @login_required
    def export_sales_excel():
        selected_month = request.args.get("month", "April")
        if selected_month not in MONTHS:
            selected_month = "April"

        selected_financial_year = selected_financial_year_from_request()
        records, columns = get_sales_invoice_details(current_user.id, selected_month, selected_financial_year)
        output = make_invoice_excel(records, columns, f"{selected_financial_year} Sales")
        filename = f"sales-invoices-{selected_financial_year}-{selected_month.lower()}.xlsx"
        return send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    @app.route("/dashboard/purchase/export/excel")
    @login_required
    def export_purchase_excel():
        selected_month = request.args.get("month", "April")
        if selected_month not in MONTHS:
            selected_month = "April"

        selected_financial_year = selected_financial_year_from_request()
        records, columns = get_purchase_invoice_details(
            current_user.id,
            selected_month,
            selected_financial_year,
        )
        output = make_invoice_excel(records, columns, f"{selected_financial_year} Purchase")
        filename = f"purchase-invoices-{selected_financial_year}-{selected_month.lower()}.xlsx"
        return send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    @app.route("/dashboard/sales/export/pdf")
    @login_required
    def export_sales_pdf():
        selected_month = request.args.get("month", "April")
        if selected_month not in MONTHS:
            selected_month = "April"

        selected_financial_year = selected_financial_year_from_request()
        records, columns = get_sales_invoice_details(current_user.id, selected_month, selected_financial_year)
        title = f"Sales Invoices - {current_user.company_name} - {selected_financial_year} - {selected_month}"
        output = make_simple_pdf(title, records, columns)
        filename = f"sales-invoices-{selected_financial_year}-{selected_month.lower()}.pdf"
        return send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype="application/pdf",
        )

    @app.route("/dashboard/purchase/export/pdf")
    @login_required
    def export_purchase_pdf():
        selected_month = request.args.get("month", "April")
        if selected_month not in MONTHS:
            selected_month = "April"

        selected_financial_year = selected_financial_year_from_request()
        records, columns = get_purchase_invoice_details(
            current_user.id,
            selected_month,
            selected_financial_year,
        )
        title = f"Purchase Invoices - {current_user.company_name} - {selected_financial_year} - {selected_month}"
        output = make_simple_pdf(title, records, columns)
        filename = f"purchase-invoices-{selected_financial_year}-{selected_month.lower()}.pdf"
        return send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype="application/pdf",
        )

    @app.route("/dashboard/documents")
    @login_required
    def documents():
        from sharepoint_docs import SharePointDocumentError, get_client_document_link

        try:
            document_link = get_client_document_link(current_user)
            configured = True
        except SharePointDocumentError as exc:
            flash(str(exc), "warning")
            document_link = ""
            configured = False

        return render_template(
            "documents.html",
            configured=configured,
            document_link=document_link,
            sharepoint_folder_path=current_user.sharepoint_folder_path or "",
        )

    @app.route("/dashboard/documents/open")
    @login_required
    def open_documents():
        from sharepoint_docs import SharePointDocumentError, get_client_document_link

        try:
            return redirect(get_client_document_link(current_user))
        except SharePointDocumentError as exc:
            flash(str(exc), "warning")
            return redirect(url_for("documents"))

    @app.cli.command("init-db")
    def init_db():
        db.create_all()
        print("Database tables created.")

    @app.cli.command("init-upload-db")
    def init_upload_db():
        init_upload_database()
        print("SQLite upload tables created.")

    @app.cli.command("create-client")
    def create_client():
        company_name = input("Company name: ").strip()
        email = input("Email: ").strip().lower()
        password = getpass("Password: ")

        if Client.query.filter_by(email=email).first():
            print("A client with that email already exists.")
            return

        client = Client(company_name=company_name, email=email)
        client.set_password(password)
        db.session.add(client)
        db.session.commit()
        print(f"Client created: {email}")

    @app.cli.command("create-admin")
    def create_admin():
        name = input("Admin name: ").strip()
        email = input("Admin email: ").strip().lower()
        password = getpass("Admin password: ")

        if Admin.query.filter_by(email=email).first():
            print("An admin with that email already exists.")
            return

        admin = Admin(name=name, email=email)
        admin.set_password(password)
        db.session.add(admin)
        db.session.commit()
        print(f"Admin created: {email}")

    @app.cli.command("reset-admin")
    def reset_admin():
        email = input("Current admin email: ").strip().lower()
        admin = Admin.query.filter_by(email=email).first()

        if admin is None:
            print("Admin not found.")
            return

        new_name = input(f"New admin name [{admin.name}]: ").strip()
        new_email = input(f"New admin email [{admin.email}]: ").strip().lower()
        new_password = getpass("New admin password: ")

        if new_name:
            admin.name = new_name

        if new_email and new_email != admin.email:
            existing_admin = Admin.query.filter(Admin.email == new_email, Admin.id != admin.id).first()
            if existing_admin:
                print("Another admin already uses that email.")
                return
            admin.email = new_email

        if new_password:
            admin.set_password(new_password)

        db.session.commit()
        print(f"Admin updated: {admin.email}")

    return app


app = create_app()


if __name__ == "__main__":
    app.run(debug=True, use_reloader=False)
